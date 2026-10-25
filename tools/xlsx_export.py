#!/usr/bin/env python3
"""
Einmaliger Export des Datenmodells aus docs/Mitgliederverwaltung_TSV_Ergoldsbach.xlsx
in CSV-Dateien unter docs/daten/ (ab dann die maßgebliche, versionierte Quelle).

Wendet dabei die freigegebenen Verbesserungen V1-V9 aus docs/PHASE0_Konzept.md an:
  V1 version/geaendert_am/geaendert_von an allen bearbeitbaren Tabs
  V2/V3 Eindeutigkeiten Ehrung und Benutzer
  V4 "halbjährlich" aus der Auswahlliste intervall entfernt
  V6 Mitglied.nur_funktionstraeger
  V7 Spalte sichtklasse im Feldkatalog
  V9 technischer Tab _Messung

Nur nötig, wenn das xlsx erneut als Ausgangspunkt dienen soll. Aufruf:
  python3 tools/xlsx_export.py
Benötigt openpyxl. Danach: npm run schema
"""
import csv, datetime, os, sys
import openpyxl

WURZEL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(WURZEL, 'docs', 'Mitgliederverwaltung_TSV_Ergoldsbach.xlsx')
ZIEL = os.path.join(WURZEL, 'docs', 'daten')
BEISPIEL = os.path.join(WURZEL, 'test', 'daten', 'beispiel')

ID_SPALTE = {}  # wird aus dem Feldkatalog (pflicht == PK) gefüllt

TAB_KLASSE = {   # Standard-Sichtklasse je Tab, Einzelfelder unten
    'Beitragskonto': 'finanz', 'Sollstellung': 'finanz', 'Beitragsgemeinschaft': 'finanz',
    'Qualifikation': 'hochsensibel',
    'Benutzer': 'verwaltung', 'Betroffenenanfrage': 'verwaltung', 'Aufbewahrungsfrist': 'verwaltung',
    'Aufnahmeantrag': 'verwaltung', '_Konfig': 'verwaltung', '_Protokoll': 'verwaltung', '_Messung': 'verwaltung',
    'Einwilligung': 'kontakt',
}
FELD_KLASSE = {
    'Mitglied': {
        'kontakt': ['strasse', 'hausnummer', 'plz', 'ort', 'land', 'geburtsdatum', 'geschlecht', 'telefon',
                    'handy', 'email', 'versandart', 'adresse_ungueltig', 'adresse_ungueltig_am',
                    'vertreter_mitglied_id', 'vertreter_name', 'vertreter_kontakt', 'bemerkung'],
        'finanz': ['beitragsart_id', 'beitragskonto_id', 'beitragsgemeinschaft_id'],
        'verwaltung': ['verarbeitung_eingeschraenkt', 'sperre_grund', 'loeschbar_ab', 'anonymisiert', 'anonymisiert_am'],
    },
    'Aufnahmeantrag': {'finanz': ['kontoinhaber_name', 'zahlweg_wunsch']},
}
TECHNISCH = ['version', 'geaendert_am', 'geaendert_von']   # Sichtklasse intern (Client braucht version)

NICHT_VERSIONIERT = ['_Konfig', '_Protokoll', '_Messung', 'Statistik_Jahr']

EINDEUTIG = {
    'Mitglied': ['mitgliedsnummer'],
    'Abteilungszugehoerigkeit': ['mitglied_id+abteilung_id'],
    'Sollstellung': ['beitragskonto_id+jahr'],
    'Ehrung': ['mitglied_id+ehrungsregel_id+funktion_id'],          # V2
    'Benutzer': ['email+rolle+abteilung_id'],                         # V3
    'Beitragskonto': ['mandatsreferenz'],
    '_Konfig': ['schluessel'],
}

STAMMDATEN_SEEDS = ['Abteilung', 'Funktionsart', 'Beitragsart', 'Qualifikationsart', 'Ehrungsregel',
                    'Aufbewahrungsfrist', '_Konfig']

KONFIG_BESCHLUSS = {   # Entscheidungen nach Phase 0
    'ehrung_zaehlweise': ('letzter_eintritt', 'Beschlossen 10/2026: ab letztem Eintritt, ein Austritt setzt die Zählung zurück. '
                          'Alternative: gesamt (alle Zeiträume), falls die Satzung geändert wird.'),
}

KONFIG_ZUSATZ = [
    ('jugend_bis_alter', '17', 'Bis zu diesem Alter (einschließlich) gilt jemand als Jugendmitglied (Sicht Jugendleiter). OFFEN: bestätigen'),
    ('snapshot_einbetten', 'ja', 'ja = Startdaten direkt in die Seite einbetten (spart einen Aufruf), nein = nachladen. Für Messvergleich.'),
    ('cache_warmhalten', 'nein', 'Zeit-Trigger, der den Snapshot-Cache vor Ablauf neu baut. Standard aus.'),
    ('messung_puffer_max', '500', 'Wie viele Messwerte maximal im Cache gepuffert werden, bevor sie nach _Messung geschrieben werden.'),
    ('sperre_wartezeit_ms', '10000', 'Wie lange ein Schreibvorgang höchstens auf die Sperre wartet.'),
    ('impressum_url', 'https://www.tsv-ergoldsbach.de/impressum', 'Link in der Fußzeile (D1). PRÜFEN'),
    ('datenschutz_url', 'https://www.tsv-ergoldsbach.de/datenschutz', 'Link in der Fußzeile (D1). PRÜFEN'),
]


def text(v):
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        if v.hour or v.minute or v.second:
            return v.strftime('%Y-%m-%dT%H:%M:%S')
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def schreibe(pfad, kopf, zeilen):
    os.makedirs(os.path.dirname(pfad), exist_ok=True)
    with open(pfad, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(kopf)
        for z in zeilen:
            w.writerow(z)


def klasse(tab, spalte):
    if spalte in TECHNISCH:
        return 'intern'
    for k, felder in FELD_KLASSE.get(tab, {}).items():
        if spalte in felder:
            return k
    return TAB_KLASSE.get(tab, 'intern')


def main():
    wb = openpyxl.load_workbook(XLSX)
    # --- Tabellenübersicht
    uebersicht = {}
    for r in wb['01_Tabellenuebersicht'].iter_rows(min_row=3, values_only=True):
        if r[0] and r[1]:
            uebersicht[r[0]] = (r[1], r[4] or '')
    # --- Feldkatalog
    felder = {}
    reihenfolge = []
    for r in wb['_Feldkatalog'].iter_rows(min_row=2, values_only=True):
        tab, gruppe, spalte, typ, pflicht, liste, beschr = r
        if not tab:
            continue
        if tab not in felder:
            felder[tab] = []
            reihenfolge.append(tab)
        felder[tab].append([spalte, typ, pflicht, liste or '', beschr or ''])
        if pflicht == 'PK':
            ID_SPALTE[tab] = spalte
    # V6
    m = felder['Mitglied']
    pos = [f[0] for f in m].index('ehrenmitglied_seit') + 1
    m.insert(pos, ['nur_funktionstraeger', 'ja/nein', 'nein', 'ja_nein',
                   'Externe Person mit Funktion (z. B. Trainer mit FZ-Pflicht) ohne Mitgliedschaft. Zählt nicht für BLSV und Beitrag. (V6)'])
    # V9
    felder['_Messung'] = [
        ['zeitpunkt', 'Zeitpunkt', 'ja', '', 'Zeitpunkt der Messung.'],
        ['funktion', 'Text', 'ja', '', 'Aufgerufene Server-Funktion.'],
        ['rolle', 'Text', 'nein', '', 'Rolle(n) des Aufrufers, keine E-Mail (Datensparsamkeit).'],
        ['dauer_ms', 'Ganzzahl', 'ja', '', 'Laufzeit auf dem Server.'],
        ['sheets_aufrufe', 'Ganzzahl', 'ja', '', 'Zahl der Zugriffe auf die Tabelle.'],
        ['cache', 'Text', 'nein', '', 'treffer / fehlgriff / leer.'],
        ['client_ms', 'Ganzzahl', 'nein', '', 'Vom Browser gemessene Gesamtzeit inkl. Übertragung, falls nachgereicht.'],
        ['ok', 'ja/nein', 'ja', 'ja_nein', 'Ob der Aufruf erfolgreich war.'],
    ]
    reihenfolge.append('_Messung')
    uebersicht['_Messung'] = ('technisch', 'Zeitmessung je Server-Aufruf (V9). Keine Personendaten.')
    # V1
    for tab in reihenfolge:
        if tab in NICHT_VERSIONIERT:
            continue
        namen = [f[0] for f in felder[tab]]
        if 'version' not in namen:
            felder[tab].append(['version', 'Ganzzahl', 'auto', '', 'Zähler für optimistische Sperre, +1 je Änderung. Nur vom Skript. (V1)'])
        if 'geaendert_am' not in namen:
            felder[tab].append(['geaendert_am', 'Zeitpunkt', 'auto', '', 'Letzte Änderung. Nur vom Skript. (V1)'])
        if 'geaendert_von' not in namen:
            felder[tab].append(['geaendert_von', 'Text', 'auto', '', 'E-Mail des Bearbeiters. Nur vom Skript. (V1)'])

    zeilen = []
    for tab in reihenfolge:
        for spalte, typ, pflicht, liste, beschr in felder[tab]:
            zeilen.append([tab, spalte, typ, pflicht, liste, klasse(tab, spalte), beschr])
    schreibe(os.path.join(ZIEL, 'feldkatalog.csv'),
             ['tabelle', 'spalte', 'typ', 'pflicht', 'auswahlliste', 'sichtklasse', 'beschreibung'], zeilen)

    tz = []
    for tab in reihenfolge:
        gruppe, zweck = uebersicht.get(tab, ('', ''))
        tz.append([tab, gruppe, ID_SPALTE.get(tab, ''), ' '.join(EINDEUTIG.get(tab, [])), zweck])
    schreibe(os.path.join(ZIEL, 'tabellen.csv'), ['tabelle', 'gruppe', 'id_spalte', 'eindeutig', 'zweck'], tz)

    # --- Auswahllisten (Langformat)
    ws = wb['_Enums']
    rows = list(ws.iter_rows(values_only=True))
    kopf = rows[0]
    ez = []
    for i, name in enumerate(kopf):
        for r in rows[1:]:
            v = r[i]
            if v is None or (isinstance(v, str) and v.startswith('Auswahllisten für')):
                continue
            if name == 'intervall' and v == 'halbjährlich':   # V4
                continue
            ez.append([name, text(v)])
        if name == 'protokoll_aktion':
            ez += [[name, 'eingesehen'], [name, 'abgewiesen']]   # Lesen sensibler Daten, abgewiesener Zugriff (A10)
    schreibe(os.path.join(ZIEL, 'auswahllisten.csv'), ['liste', 'wert'], ez)

    # --- Startbelegung (Stammdaten) und Beispieldaten (alles erfunden)
    for tab in reihenfolge:
        if tab == '_Messung':
            continue
        ws = wb[tab]
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        kopf = list(rows[0])
        daten = [[text(c) for c in r[:len(kopf)]] for r in rows[1:]]
        if tab == '_Konfig':
            for z in daten:
                if z[0] in KONFIG_BESCHLUSS:
                    z[1], z[2] = KONFIG_BESCHLUSS[z[0]]
            daten += [list(z) for z in KONFIG_ZUSATZ]
        ziel = os.path.join(ZIEL, 'startbelegung', tab + '.csv') if tab in STAMMDATEN_SEEDS \
            else os.path.join(BEISPIEL, tab + '.csv')
        schreibe(ziel, kopf, daten)
    print('Export fertig:', ZIEL)


if __name__ == '__main__':
    sys.exit(main())
